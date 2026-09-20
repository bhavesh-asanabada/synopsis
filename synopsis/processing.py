"""Single-worker document pipeline: retain originals, extract all pages, OCR as needed."""
import logging
import os
from threading import Lock
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pypdfium2 as pdfium
import pytesseract
from PIL import Image, ImageOps, ImageSequence
from pypdf import PdfReader
from docx import Document

from .metadata import infer_metadata, lookup_doi
from .documents import ocr_page, pdf_geometry, detect_language, reference_candidates, fingerprint

log = logging.getLogger(__name__)
ALLOWED = {'.pdf', '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.webp', '.txt', '.md', '.docx'}


class Processor:
    def __init__(self, store):
        self.store = store
        # PDFium is not thread-safe. Serialize PDF rendering and OCR jobs.
        self.pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix='synopsis-ocr')
        self.pending = set()
        self.lock = Lock()
        self.cloud = None

    def submit(self, item_id):
        if self.cloud:
            self.cloud.enqueue(item_id)
        with self.lock:
            if item_id in self.pending:
                return
            self.pending.add(item_id)
        future = self.pool.submit(self.process, item_id)
        def finished(_):
            with self.lock:
                self.pending.discard(item_id)
        future.add_done_callback(finished)

    def process(self, item_id):
        item = self.store.get(item_id)
        if not item:
            return
        text_parts, embedded, warnings, page_records, tables = [], {}, [], [], []
        used_ocr = False
        page_count = 0
        try:
            self.store.update(item_id, {'status': 'processing', 'progress': 'Reading document…', 'warnings': []})
            path = self.store.document_path(item)
            suffix = path.suffix.lower()
            language = self.store.setting('ocrLanguage', 'eng')

            def recognize(image):
                # Bound raster size while keeping enough resolution for ordinary scans.
                image = ImageOps.exif_transpose(image).convert('RGB')
                image.thumbnail((4000, 4000))
                return ocr_page(image, language)

            if suffix == '.pdf':
                reader = PdfReader(path)
                if reader.is_encrypted and not reader.decrypt(''):
                    raise ValueError('This PDF is password-protected. Upload an unlocked copy.')
                meta = reader.metadata
                embedded = {'title': str(meta.title or ''), 'author': str(meta.author or '')} if meta else {}
                page_count = len(reader.pages)
                with pdfium.PdfDocument(str(path)) as document:
                    for number, page in enumerate(reader.pages):
                        self.store.update(item_id, {'progress': f'Reading page {number + 1} of {page_count}…'})
                        text = page.extract_text() or ''
                        if len(''.join(text.split())) < 40:
                            self.store.update(item_id, {'progress': f'OCR · page {number + 1} of {page_count}…'})
                            rendered = document[number]
                            try:
                                width, height = rendered.get_size()
                                scale = min(2.5, 4000 / max(width, height))
                                bitmap = rendered.render(scale=scale)
                                try:
                                    record = recognize(bitmap.to_pil())
                                    text = record['text']
                                finally:
                                    bitmap.close()
                            finally:
                                rendered.close()
                            used_ocr = True
                        else:
                            record = {'text':text, 'method':'embedded-text', 'confidence':None, 'rotation':0}
                        page_records.append({**record, 'page':number+1})
                        text_parts.append(text)
                geometry, tables = pdf_geometry(path)
                for record, geo in zip(page_records, geometry):
                    if record['method'] != 'ocr':
                        record.update(geo)
            elif suffix in {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.webp'}:
                with Image.open(path) as image:
                    page_count = getattr(image, 'n_frames', 1)
                    for index, frame in enumerate(ImageSequence.Iterator(image)):
                        self.store.update(item_id, {'progress': f'OCR · page {index + 1} of {page_count}…'})
                        record = recognize(frame.copy())
                        page_records.append({**record, 'page':index+1})
                        text_parts.append(record['text'])
                used_ocr = True
            elif suffix == '.docx':
                document = Document(path)
                embedded = {'title': document.core_properties.title or '', 'author': document.core_properties.author or ''}
                text_parts = ['\n'.join(p.text for p in document.paragraphs)]
                text_parts.extend('\n'.join('\t'.join(c.text for c in row.cells) for row in t.rows) for t in document.tables)
                tables = [{'page':None, 'rows':[[c.text for c in row.cells] for row in t.rows], 'method':'docx-table', 'reviewed':False} for t in document.tables]
            else:
                text_parts = [path.read_text(encoding='utf-8-sig')]

            text = '\n\n'.join(text_parts)
            if not page_records:
                page_records = [{'page':1, 'text':text, 'method':'document-text', 'confidence':None, 'words':[], 'pagination':'logical'}]
            for page in page_records:
                page['hash'] = fingerprint(page['text'])
                if page.get('confidence') is not None and page['confidence'] < .75:
                    warnings.append(f"Page {page['page']}: OCR confidence is low; verify the original scan.")
            metadata = infer_metadata(text, item['fileName'], embedded)
            opening=text[:12000].lower()
            metadata['type'] = ('thesis' if any(t in opening for t in ('doctoral thesis','doctoral dissertation','master\'s thesis'))
                                else 'report' if 'technical report' in opening else 'paper-conference' if 'conference proceedings' in opening
                                else 'article-journal' if 'abstract' in opening else 'document')
            provenance = {k:{'source':'embedded-metadata' if k in ('title','authors') and embedded.get('title' if k=='title' else 'author') else 'text-heuristic',
                             'reviewRequired':True, 'value':v} for k,v in metadata.items() if v}
            for key, origin in provenance.items():
                value=origin['value']
                needle=value if isinstance(value,str) else ''
                match=next((p for p in page_records if needle and needle in p['text']),None)
                if match:
                    origin.update(page=match['page'],excerpt=needle[:300])
            source = 'ocr' if used_ocr else 'document'
            if metadata['doi'] and self.store.setting('metadataLookup', True):
                self.store.update(item_id, {'progress': 'Looking up bibliographic metadata…'})
                try:
                    found = {k:v for k,v in lookup_doi(metadata['doi']).items() if v}
                    provenance.update({k:{'source':'crossref', 'doi':metadata['doi'], 'reviewRequired':True, 'value':v} for k,v in found.items()})
                    metadata.update(found)
                    source = 'crossref'
                except Exception:
                    warnings.append('DOI lookup was unavailable. Review the metadata extracted from your document.')
            if not text.strip():
                warnings.append('No readable text was found. The original document is saved; enter metadata manually.')
            if not metadata.get('authors'):
                warnings.append('Authors could not be identified reliably. Please add them manually.')
            # Preserve any edits made while the job was running.
            self.store.update(item_id, lambda current: {
                                       **{k: v for k, v in metadata.items() if k not in current.get('editedFields', [])},
                                       'text': text, 'pageCount': page_count, 'source': source,
                                       'pageTexts': page_records, 'tables':tables, 'references':reference_candidates(text),
                                       'language':detect_language(text), 'extractionVersion':2,
                                       'provenance':{**provenance, **{k:v for k,v in current.get('provenance',{}).items() if k in current.get('editedFields',[])}},
                                       'ocrUsed': used_ocr, 'warnings': warnings, 'status': 'processing',
                                       'reviewed': False, 'progress': 'Applying organization rules…'})
            from .services import apply_rules
            apply_rules(self.store, self.store.get(item_id))
            self.store.update(item_id,{'status':'ready','progress':'Ready for review'})
        except Exception as exc:
            log.exception('Document processing failed for %s', item_id)
            if isinstance(exc, pytesseract.TesseractNotFoundError):
                message = 'Tesseract is not installed. Install it and retry OCR; your original file is saved.'
            else:
                message = f'Could not process this document: {str(exc)[:300]}'
            self.store.update(item_id, {'status': 'error', 'progress': message, 'warnings': [message],
                                       'text': '\n\n'.join(text_parts), 'reviewed': False})
